"""Generate 5 test CVs as real DOCX files in tests/cvs/."""
import os

import docx

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "cvs")

CVS = {
    "01_green_priya_raman.docx": [
        ("h", "Priya Raman"),
        ("p", "Senior Machine Learning Engineer · priya.raman@example.com · +44 7700 900123 · London"),
        ("h2", "Summary"),
        ("p", "Senior ML engineer with 9 years of Python experience, specialising in computer vision. "
              "I have designed, trained and deployed CV models serving millions of requests per day."),
        ("h2", "Experience"),
        ("p", "Senior ML Engineer, Visionly Ltd - Mar 2021 to present"),
        ("b", "Lead a team of 4 building real-time object detection (YOLOv8, PyTorch) for retail shelf analytics."),
        ("b", "Deployed models to production on Kubernetes with Triton; p95 latency 40 ms at 3k req/s."),
        ("b", "Built MLOps pipeline (MLflow, Airflow) for automated retraining and drift monitoring."),
        ("p", "Machine Learning Engineer, DeepRoad Autonomous - Jun 2017 to Feb 2021"),
        ("b", "Developed semantic segmentation models for lane detection in Python/PyTorch; shipped to vehicle fleet."),
        ("b", "Optimised inference with TensorRT, reducing GPU cost by 35%."),
        ("p", "Software Engineer (Python), DataNest - Sep 2015 to May 2017"),
        ("b", "Python backend services and image-processing pipelines with OpenCV."),
        ("h2", "Education"),
        ("p", "MSc Computer Vision, University of Edinburgh, 2015. BSc Computer Science, 2014."),
        ("h2", "Skills"),
        ("p", "Python, PyTorch, TensorFlow, OpenCV, Kubernetes, Docker, AWS, MLflow, Triton, TensorRT"),
    ],
    "02_gray_tom_becker.docx": [
        ("h", "Tom Becker"),
        ("p", "Data Scientist · tom.becker@example.com · +49 151 2345678"),
        ("h2", "Summary"),
        ("p", "Data scientist with Python experience and some exposure to machine learning and image analysis."),
        ("h2", "Experience"),
        ("p", "Data Scientist, RetailMetrics GmbH - Jan 2023 to present"),
        ("b", "Python analytics and dashboards; trained a scikit-learn churn model that was rolled into the monthly reporting pipeline."),
        ("b", "Fine-tuned a pretrained ResNet in PyTorch to classify product photos; piloted in two stores."),
        ("p", "Career break - no details given"),
        ("p", "Data Analyst, InsureCo - Jul 2016 to Nov 2020"),
        ("b", "SQL reporting, Python scripts for data cleaning, some forecasting models."),
        ("h2", "Education"),
        ("p", "BSc Mathematics, TU Munich, 2016"),
        ("h2", "Skills"),
        ("p", "Python, pandas, scikit-learn, SQL, Tableau, basic PyTorch"),
    ],
    "03_red_gary_doyle.docx": [
        ("h", "Gary Doyle"),
        ("p", "Construction Project Manager · gary.doyle@example.com · +1 555 0142"),
        ("h2", "Summary"),
        ("p", "Construction project manager with 15 years delivering commercial building projects on time and budget."),
        ("h2", "Experience"),
        ("p", "Senior Project Manager, BuildRight Construction - 2014 to present"),
        ("b", "Managed $40M office and retail builds; coordinated subcontractors and site safety."),
        ("p", "Site Supervisor, Doyle & Sons Builders - 2009 to 2014"),
        ("b", "Supervised crews of 20+, scheduling, procurement and inspections."),
        ("h2", "Education"),
        ("p", "BSc Construction Management, Ohio State University, 2008"),
        ("h2", "Skills"),
        ("p", "MS Project, Procore, budgeting, OSHA compliance, stakeholder management"),
    ],
    "04_contradiction_alex_kim.docx": [
        ("h", "Alex Kim"),
        ("p", "Lead Machine Learning Engineer · alex.kim@example.com · +1 555 0199"),
        ("h2", "Summary"),
        ("p", "Lead ML engineer with 12 years experience in Python, computer vision and deploying models to production."),
        ("h2", "Experience"),
        ("p", "Lead ML Engineer, PixelForge AI - Aug 2022 to present"),
        ("b", "Lead computer vision work: face detection and image classification in PyTorch, deployed on AWS SageMaker."),
        ("p", "ML Engineer, SnapSense - Jul 2020 to Jul 2022"),
        ("b", "Trained and shipped CNN models for image tagging into production APIs (Python, FastAPI)."),
        ("h2", "Education"),
        ("p", "BSc Computer Science, University of Toronto - graduated 2020"),
        ("h2", "Skills"),
        ("p", "Python, PyTorch, OpenCV, AWS, Docker, FastAPI"),
    ],
    "05_duplicate_priya_raman.docx": [
        ("h", "Priya Raman"),
        ("p", "ML Engineer (Computer Vision) · Priya.Raman@Example.com · +44 7700 900123"),
        ("h2", "Profile"),
        ("p", "Senior ML engineer, 9 years Python, computer vision specialist with production deployments."),
        ("h2", "Experience"),
        ("p", "Senior ML Engineer, Visionly Ltd - Mar 2021 to present: real-time object detection in PyTorch, "
              "deployed to production on Kubernetes."),
        ("p", "Machine Learning Engineer, DeepRoad Autonomous - Jun 2017 to Feb 2021: lane segmentation models shipped to fleet."),
        ("p", "Software Engineer (Python), DataNest - Sep 2015 to May 2017."),
        ("h2", "Education"),
        ("p", "MSc Computer Vision, University of Edinburgh, 2015."),
    ],
}


def main():
    os.makedirs(OUT, exist_ok=True)
    for name, blocks in CVS.items():
        d = docx.Document()
        for kind, text in blocks:
            if kind == "h":
                d.add_heading(text, level=0)
            elif kind == "h2":
                d.add_heading(text, level=2)
            elif kind == "b":
                d.add_paragraph(text, style="List Bullet")
            else:
                d.add_paragraph(text)
        d.save(os.path.join(OUT, name))
        print("wrote", os.path.join("tests", "cvs", name))


if __name__ == "__main__":
    main()
