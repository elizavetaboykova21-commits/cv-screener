"""End-to-end test against the RUNNING app: upload -> screen -> poll -> check.

Usage: python tests/e2e.py [http://localhost:5153]
"""
import glob
import json
import os
import sys
import time
import uuid

import requests

ARGS = [a for a in sys.argv[1:] if not a.startswith("--")]
BASE = ARGS[0] if ARGS else "http://localhost:5153"
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)

SPEC = {
    "job_description": "Senior Machine Learning Engineer to design, train and deploy computer-vision "
                       "models (detection, segmentation, classification) for our retail analytics "
                       "product. You will own models from research to production.",
    "must_haves": "5+ years of professional Python experience\n"
                  "Hands-on computer vision / deep learning experience\n"
                  "Has deployed ML models to production",
    "nice_to_haves": "PyTorch\nMLOps (MLflow, Kubeflow, Airflow)\nKubernetes",
    "red_flags": "Unexplained employment gaps over 12 months\nClaims that don't match dated history",
}
EXPECT = {1: "GREEN", 2: "GRAY", 3: "RED"}
RUN_AGENT = "--no-agent" not in sys.argv


def main():
    paths = sorted(glob.glob(os.path.join(HERE, "cvs", "*.docx")))
    job_id = str(uuid.uuid4())
    files = [("files", (os.path.basename(p), open(p, "rb"))) for p in paths]
    up = requests.post(BASE + "/api/upload", data={"job_id": job_id}, files=files, timeout=60)
    print("UPLOAD", up.status_code)
    print(json.dumps(up.json(), indent=2, ensure_ascii=False))
    up.raise_for_status()

    sc = requests.post(BASE + "/api/screen", json={"job_id": job_id, **SPEC, "run_agent": RUN_AGENT}, timeout=30)
    print("SCREEN", sc.status_code, sc.text.strip())
    if sc.status_code != 200:
        sys.exit(2)

    t0, last = time.time(), None
    while True:
        st = requests.get(BASE + "/api/status/" + job_id, timeout=30).json()
        cur = (st["status"], st["progress"], st.get("agent_progress"))
        if cur != last:
            if st["status"] == "agent":
                print("  agent    %d / %d  (%.0fs)" % (st["agent_progress"], st["agent_total"], time.time() - t0))
            else:
                print("  progress %d / %d  (%.0fs)" % (st["progress"], st["total"], time.time() - t0))
            last = cur
        if st["status"] == "done":
            break
        time.sleep(1.2)

    print("\nRESULTS")
    for f in st["files"]:
        r = f["result"] or {}
        print("#%d %-34s %-18s dedup=%-5s name=%r email=%r" % (
            f["index"], f["filename"], r.get("verdict"), f.get("deduped"), r.get("name"), r.get("email")))
        if r.get("error"):
            print("     error:", r["error"])
        for q in r.get("question_marks", []):
            print("     ?", q)
        a = f.get("agent") or {}
        if a.get("skipped"):
            print("     AGENT skipped:", a["skipped"])
        for x in a.get("actions", []):
            print("     AGENT %d. %s | why: %s" % (x["step"], x["tool"], x["reason"]))
            print("        " + x["output"].replace("\n", "\n        "))
        if a.get("error"):
            print("     AGENT error:", a["error"])
    print("sheet:", st["sheet"], "| deduped:", st.get("deduped"), "| run error:", st.get("error"))

    by = {f["index"]: f for f in st["files"]}
    checks = []
    for i, want in EXPECT.items():
        got = (by[i]["result"] or {}).get("verdict")
        checks.append(("CV %d expected %s" % (i, want), got == want, got))
    qm = " ".join(by[4]["result"].get("question_marks", [])).lower()
    checks.append(("CV 4 flags 12-years vs 2020 graduation",
                   "12" in qm and "2020" in qm, by[4]["result"].get("question_marks")))
    checks.append(("CV 5 marked duplicate of CV 1",
                   by[5].get("deduped") is True and by[5].get("duplicate_of") == 1, by[5].get("duplicate_of")))
    xlsx = os.path.join(ROOT, st["sheet"]["path"]) if st["sheet"] else ""
    checks.append(("screening.xlsx exists", bool(xlsx) and os.path.exists(xlsx), xlsx))
    if xlsx and os.path.exists(xlsx):
        from openpyxl import load_workbook
        ws = load_workbook(xlsx).active
        rows = [(ws.cell(r, 3).value, ws.cell(r, 1).fill.fgColor.rgb) for r in range(2, ws.max_row + 1)]
        print("xlsx headers:", [c.value for c in ws[1]][:10])
        checks.append(("xlsx has Agent action / Output columns",
                       [ws.cell(1, 9).value, ws.cell(1, 10).value] == ["Agent action", "Output"], None))
        print("xlsx rows (verdict, fill):", rows)
        order = {"GREEN": 0, "GRAY": 1, "RED": 2, "PROCESSING FAILED": 3}
        keys = [order[v.replace(" (dup)", "")] for v, _ in rows]
        checks.append(("xlsx sorted GREEN>GRAY>RED>FAILED", keys == sorted(keys), keys))
        fills = {"GREEN": "C6EFCE", "GRAY": "FFE699", "RED": "FFC7CE", "PROCESSING FAILED": "D9D9D9"}
        checks.append(("xlsx row colours match verdict",
                       all(c.endswith(fills[v.replace(" (dup)", "")]) for v, c in rows), None))

    if RUN_AGENT:
        eligible = [f for f in st["files"] if not f.get("deduped") and (f["result"] or {}).get("verdict") != "PROCESSING FAILED"]
        checks.append(("agent acted on every unique candidate without errors",
                       all((f.get("agent") or {}).get("actions") and not f["agent"].get("error") for f in eligible),
                       [(f["index"], (f.get("agent") or {}).get("error")) for f in eligible]))
        checks.append(("agent at most 3 tool calls each",
                       all(len((f.get("agent") or {}).get("actions", [])) <= 3 for f in st["files"]), None))
        checks.append(("duplicate skipped by agent", (by[5].get("agent") or {}).get("skipped") == "duplicate", by[5].get("agent")))
        log = os.path.join(ROOT, os.path.dirname(st["sheet"]["path"]), "agent_log.jsonl") if st["sheet"] else ""
        checks.append(("agent_log.jsonl written", bool(log) and os.path.exists(log), log))
    print("\nCHECKS")
    for name, ok, info in checks:
        print("  %s %s%s" % ("PASS" if ok else "FAIL", name, "" if ok else "  -> %r" % (info,)))
    sys.exit(0 if all(ok for _, ok, _ in checks) else 1)


if __name__ == "__main__":
    main()
