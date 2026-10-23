"""Text extraction for PDF / DOCX / TXT CVs."""
import os

MIN_PDF_CHARS = 200
SUPPORTED = {".pdf", ".docx", ".txt"}


class ExtractionError(Exception):
    pass


def _pdf(path):
    from pdfminer.high_level import extract_text

    try:
        text = extract_text(path) or ""
    except Exception as e:  # corrupt / encrypted PDFs
        raise ExtractionError("unreadable PDF (%s)" % type(e).__name__)
    if len(text.strip()) < MIN_PDF_CHARS:
        raise ExtractionError(
            "scanned or image-only PDF (%d chars of text) - needs OCR" % len(text.strip())
        )
    return text


def _docx(path):
    import docx

    try:
        d = docx.Document(path)
    except Exception as e:
        raise ExtractionError("unreadable DOCX (%s)" % type(e).__name__)
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


def _txt(path):
    with open(path, "rb") as f:
        raw = f.read()
    for enc in ("utf-8", "utf-16", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ExtractionError("could not decode text file")


def extract_text(path):
    """Return extracted text, or raise ExtractionError with a human reason."""
    ext = os.path.splitext(path)[1].lower()
    if ext not in SUPPORTED:
        raise ExtractionError("unsupported file type %s" % (ext or "(none)"))
    text = {".pdf": _pdf, ".docx": _docx, ".txt": _txt}[ext](path)
    if not text.strip():
        raise ExtractionError("no text found")
    return text
