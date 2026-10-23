"""Build the per-run Excel results sheet."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ORDER = {"GREEN": 0, "GRAY": 1, "RED": 2, "PROCESSING FAILED": 3}
FILLS = {
    "GREEN": "C6EFCE",
    "GRAY": "FFE699",  # amber: "maybe, verify"
    "RED": "FFC7CE",
    "PROCESSING FAILED": "D9D9D9",
}
HEADERS = [
    "Name", "Email", "Verdict", "Current Title", "Current Company",
    "Why Fit", "Why Not Fit", "Question Marks", "Agent action", "Output",
]
NOTE_COLS = 7  # K..Q left blank for recruiter notes
WIDTHS = [24, 30, 18, 26, 24, 60, 60, 70, 30, 90]


def sort_key(f):
    r = f.get("result") or {}
    return (ORDER.get(r.get("verdict"), 3), f.get("index", 0))


def agent_cells(f):
    """(Agent action, Output) text for one candidate."""
    a = f.get("agent")
    if not a:
        return "", ""
    if a.get("skipped"):
        return "skipped (%s)" % a["skipped"], ""
    acts = a.get("actions") or []
    action = ", ".join(x["label"] for x in acts)
    blocks = ["[%s]\nWhy: %s\n%s" % (x["label"], x["reason"], x["output"]) for x in acts]
    if a.get("error"):
        action = (action + "; " if action else "") + "agent error"
        blocks.append("[agent error] " + a["error"])
    return action, "\n\n".join(blocks)


def build_sheet(files, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Screening"

    headers = HEADERS + ["Notes %d" % i for i in range(1, NOTE_COLS + 1)]
    ws.append(headers)
    head_fill = PatternFill("solid", fgColor="1F2937")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(vertical="center")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(top=thin, bottom=thin, left=thin, right=thin)

    for f in sorted(files, key=sort_key):
        r = f.get("result") or {}
        verdict = r.get("verdict", "PROCESSING FAILED")
        why_not = r.get("why_not_fit", "")
        qm = list(r.get("question_marks") or [])
        if f.get("deduped"):
            qm.insert(0, "DUPLICATE of %s" % f.get("duplicate_of_name", "an earlier CV"))
        if verdict == "PROCESSING FAILED":
            why_not = r.get("error") or f.get("error") or "processing failed"
        row = [
            r.get("name") or f.get("filename", ""),
            r.get("email", ""),
            verdict + (" (dup)" if f.get("deduped") else ""),
            r.get("current_title", ""),
            r.get("current_company", ""),
            r.get("why_fit", ""),
            why_not,
            "; ".join(qm),
            *agent_cells(f),
        ]
        ws.append(row + [""] * NOTE_COLS)
        fill = PatternFill("solid", fgColor=FILLS.get(verdict, "D9D9D9"))
        for c in ws[ws.max_row][: len(HEADERS)]:
            c.fill = fill
            c.border = border
            c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(ws.max_row, 3).font = Font(bold=True)

    for i, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for i in range(len(HEADERS) + 1, len(HEADERS) + 1 + NOTE_COLS):
        ws.column_dimensions[get_column_letter(i)].width = 18
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(HEADERS)), ws.max_row)
    wb.save(path)
    return path
